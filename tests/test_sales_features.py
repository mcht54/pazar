"""Satış puanı, fırsat kartları, çözüm rehberi, satış notu, CRM, dışa aktarma, dashboard ve tekrar tarama davranışları."""

import csv
import io
import re
from pathlib import Path

import pytest

from packages.crm import CRM_STAGES
from packages.db.models import Business, Region, Sector
from services.knowledge.guides import CATEGORIES, CHECK_TO_GUIDE, GUIDES, NO_GUIDE_CHECKS, guide_for_check, render_guide
from services.research.dedupe import same_business
from services.rule_engine.checks import build_gbp_checks, build_social_checks, build_website_checks
from services.rule_engine.opportunities import build_opportunities
from services.rule_engine.sales import assess
from services.rule_engine.scoring import ASSUMPTION_CAP, AREA_CAP, compute_score

from tests.test_rule_engine import ctx, good_signals, google_business, listing_business
from services.integrations.website_crawler.base import WebsiteSignals

ROOT = Path(__file__).resolve().parent.parent


def _analyze(context, website, signals, gbp_biz):
    checks = build_website_checks(context, website, signals) + build_gbp_checks(context, gbp_biz)
    a = assess(context, checks, has_website=bool(website), contactable_phone=True, contactable_email=False)
    return checks, a


# ================================================================ PUAN
def test_score_is_high_when_no_website_and_low_for_strong_presence():
    weak, a1 = _analyze(ctx(source="google_places"), None, None, google_business(website=None, photo_count=0, google_review_count=2))
    strong, a2 = _analyze(ctx(source="google_places"), "https://kaya.example", good_signals(), google_business())
    s_weak, s_strong = compute_score(weak, a1.possible_services), compute_score(strong, a2.possible_services)
    assert s_weak["score"] >= 45 and s_weak["band"] in ("Yüksek fırsat", "Çok yüksek fırsat")
    assert s_strong["score"] < 20, "güçlü dijital varlığı olan işletme yüksek puan almamalı"
    assert s_weak["score"] > s_strong["score"] + 30


def test_score_is_deterministic_bounded_and_explained():
    checks, a = _analyze(ctx(source="google_places"), None, None, google_business(website=None, photo_count=0, google_review_count=1, google_rating=3.2))
    r1, r2 = compute_score(checks, a.possible_services), compute_score(checks, a.possible_services)
    assert r1 == r2, "aynı girdi aynı puanı vermeli (rastgelelik yok)"
    assert 0 <= r1["score"] <= 100
    assert r1["items"], "puanın nedenleri (kalemleri) listelenmeli"
    assert all(i["kind"] == "tespit" and i["points"] > 0 and i["explanation"] and i["label"] for i in r1["items"])
    total = sum(i["points"] for i in r1["items"]) + sum(x["points"] for x in r1["adjustments"]) + min(sum(x["points"] for x in r1["assumptions"]), ASSUMPTION_CAP)
    assert r1["score"] == min(100, max(0, total)), "puan kalemlerinin toplamıyla birebir eşleşmeli"


def test_unknown_checks_do_not_add_points():
    checks = build_gbp_checks(ctx(source="google_maps"), listing_business(phone="1"))
    assert all(c.status == "unknown" for c in checks)
    assert compute_score(checks, [])["score"] == 0


def test_no_website_gives_larger_impact_than_a_minor_gap():
    no_site, _ = _analyze(ctx(source="google_places"), None, None, google_business())
    minor_signals = good_signals(has_blog=False)
    minor, _ = _analyze(ctx(source="google_places"), "https://kaya.example", minor_signals, google_business())
    assert compute_score(no_site, [])["score"] > compute_score(minor, [])["score"] + 20


def test_assumption_points_are_capped_and_labeled_unverified():
    poss = [{"service": s, "basis": "x"} for s in ("Google Ads", "Sosyal Medya Reklamları", "Tabela", "Matbaa", "Davetiye", "Araç Giydirme", "Fotoğraf ve Video İçerik")]
    r = compute_score([], poss)
    assert r["score"] <= ASSUMPTION_CAP
    assert all("DOĞRULANMADI" in x["label"] and x["kind"] == "varsayım" for x in r["assumptions"])


def test_area_cap_prevents_many_small_findings_from_inflating():
    s = good_signals(has_blog=False, has_about_page=False, has_references_page=False, sitemap_present=False, schema_types=[], social_links={}, copyright_year=2019,
                     meta_description=None, h1_texts=[], title="Ana Sayfa", whatsapp_link_present=False, maps_link_present=False)
    checks, a = _analyze(ctx(source="google_places"), "https://kaya.example", s, google_business())
    r = compute_score(checks, [])
    web_points = sum(i["points"] for i in r["items"] if i["area"] == "website") + sum(x["points"] for x in r["adjustments"] if "Web sitesi" in x["label"])
    assert web_points <= AREA_CAP["website"]


# ================================================================ FIRSAT KARTLARI
def test_opportunities_separate_verified_from_possible_and_never_invent_problems():
    checks, a = _analyze(ctx(source="google_places"), None, None, google_business(website=None))
    ops = build_opportunities(a)
    assert ops["evidence"], "web sitesi yok → tespite dayalı fırsat olmalı"
    first = ops["evidence"][0]
    assert first["service"] == "Kurumsal Web Sitesi" and first["level"] == "Yüksek" and first["verified"] is True
    assert first["problems"][0]["label"] == "Web sitesi bulunamadı" and first["problems"][0]["evidence"]
    assert first["why"] and first["rationale"] and first["offer"]
    for p in ops["possible"]:
        assert p["verified"] is False and p["level"] == "Doğrulanmadı" and p["problems"] == []
        assert "Tespit edilmedi" in p["problem_text"], "kanıtsız hizmette sorun uydurulmamalı"
    assert not ({o["service"] for o in ops["evidence"]} & {p["service"] for p in ops["possible"]})


def test_healthy_business_has_no_verified_opportunities():
    _, a = _analyze(ctx(source="google_places"), "https://kaya.example", good_signals(), google_business())
    assert build_opportunities(a)["evidence"] == []


def test_possible_services_are_sector_based_for_davetiye_and_arac_giydirme():
    from services.rule_engine.sector_profiles import get_sector_profile

    wedding = ctx("Düğün ve Davet Salonu", "Profesyonel Hizmetler", phrases=["düğün salonu"])
    wedding.sector_group = "Profesyonel Hizmetler"
    _, a = _analyze(wedding, "https://kaya.example", good_signals(), google_business())
    assert "Davetiye" in {p["service"] for p in a.possible_services}
    mover = ctx("Nakliyat", "Ulaşım ve Lojistik", phrases=["nakliyat"])
    mover.sector_group = "Ulaşım ve Lojistik"
    _, a2 = _analyze(mover, "https://kaya.example", good_signals(), google_business())
    assert "Araç Giydirme" in {p["service"] for p in a2.possible_services}


# ================================================================ REHBERLER
def test_every_analysis_check_key_has_a_guide():
    keys = set(re.findall(r'key="([a-z_]+)"', (ROOT / "services/rule_engine/checks.py").read_text(encoding="utf-8")))
    social_keys = {"social_presence"}
    unmapped = []
    for key in keys:
        candidates = [f"web.{key}", f"gbp.{key}", f"social.{key}"]
        if not any(c in CHECK_TO_GUIDE for c in candidates) and not any(c in NO_GUIDE_CHECKS for c in candidates) and not key.startswith("social_"):
            unmapped.append(key)
    assert not unmapped, f"rehberi olmayan kontrol anahtarları: {unmapped}"
    assert social_keys <= {k.split('.', 1)[1] for k in CHECK_TO_GUIDE}


def test_guides_have_all_eight_sections_and_are_turkish():
    assert len(GUIDES) >= 50
    assert set(CATEGORIES) == {g.category for g in GUIDES.values()}
    english = re.compile(r"\b(the|and|please|click|website|error|not found)\b", re.IGNORECASE)
    for guide in GUIDES.values():
        rendered = render_guide(guide)["sections"]
        assert rendered["what"] and rendered["why"] and rendered["solution"] and rendered["pitch"] and rendered["offer"], guide.id
        assert len(rendered["steps"]) >= 3 and len(rendered["questions"]) >= 1 and len(rendered["verify"]) >= 1, guide.id
        assert guide.services, f"{guide.id}: hizmet eşleşmesi yok"
        text = " ".join([rendered["what"], rendered["why"], rendered["solution"], rendered["pitch"], *rendered["steps"], *rendered["questions"], *rendered["verify"]])
        assert not english.search(text), f"{guide.id}: İngilizce ifade"
        assert "{" not in text, f"{guide.id}: doldurulmamış yer tutucu"


def test_required_library_topics_exist():
    from packages.localization import fold

    titles = " | ".join(fold(g.title) for g in GUIDES.values())
    for topic in ["web sitesi yok", "https", "mobil", "title", "meta description", "h1", "seo içerik", "iletişim", "whatsapp", "local seo", "schema",
                  "kategori", "hizmet", "açıklama", "fotoğraf", "yorum sayısı", "cevap verilmemiş", "çalışma saatleri", "web sitesi bağlantısı",
                  "aktif değil", "içerik düzensiz", "profil eksikleri", "iletişim eksikleri", "google ads", "yerel arama", "marka görünürlüğü",
                  "dönüşüm takibi", "kartvizit", "broşür", "menü", "tabela", "branda", "promosyon", "davetiye", "kurumsal materyaller"]:
        assert fold(topic) in titles, f"Çözüm Rehberi'nde konu yok: {topic}"


def test_guide_wording_never_states_uncertain_effects_as_facts():
    for guide in GUIDES.values():
        assert not re.search(r"(kesinlikle|garanti|%\s?\d+ artar|mutlaka artar)", guide.why + guide.pitch, re.IGNORECASE), guide.id


def test_render_guide_uses_real_business_context_and_generic_fallback():
    g = GUIDES["web_yok"]
    real = render_guide(g, business="Si-Ser İşitme", place="Adapazarı", sector="İşitme Cihazı Merkezi", phrase="işitme cihazı")
    assert "Si-Ser İşitme" in real["sections"]["pitch"] and "Adapazarı" in real["sections"]["pitch"]
    generic = render_guide(g)
    assert "Si-Ser" not in generic["sections"]["pitch"] and "işletmeniz" in " ".join([generic["sections"]["pitch"], generic["sections"]["what"]])


def test_guide_for_check_maps_area_names():
    assert guide_for_check("website", "presence").id == "web_yok"
    assert guide_for_check("gbp", "review_count").id == "gbp_yorum_az"
    assert guide_for_check("gbp", "website").id == "gbp_web"
    assert guide_for_check("social", "social_presence").id == "sosyal_yok"


# ================================================================ DEDUPE
def test_same_business_rules_are_conservative_about_branches():
    a = {"name": "Si-Ser İşitme Cihazları", "phone": "0264 277 45 05", "lat": 40.7582843, "lng": 30.3877655, "place_id": "0x1"}
    assert same_business(a, {**a, "name": "Si-Ser Adapazarı İşitme Cihazları", "lat": 40.7582889, "lng": 30.3877544, "place_id": "0x2"})[0]
    assert same_business(a, {**a, "place_id": "0x1", "lat": 41.0, "lng": 29.0})[0], "aynı Google kaydı her koşulda aynı"
    branch = {"name": "Sakarya İşitme Cihazları - Şube", "phone": "0545 485 37 36", "lat": 40.7600, "lng": 30.3900}
    main = {"name": "Sakarya İşitme Cihazları", "phone": "0545 485 37 36", "lat": 40.7583, "lng": 30.3840}
    assert not same_business(branch, main)[0], "aynı telefonu paylaşan ama farklı konumdaki şubeler ayrı işletmedir"
    assert not same_business({"name": "İdis İşitme", "lat": 1.0, "lng": 1.0}, {"name": "Adaduy İşitme", "lat": 1.0, "lng": 1.0})[0]


# ================================================================ API: CRM / rehber / not / dışa aktarma / panel
def _analyzed_business(client, db, name="Test Kliniği", phone="0532 111 22 33", website=None, category=None, sector_name="İşitme Cihazı Merkezi", user_id=None):
    from packages.db.models import AnalysisJob
    from services.worker.tasks.analysis import run_analysis_job

    region = db.query(Region).filter_by(name="Serdivan").one()
    sector = db.query(Sector).filter_by(name=sector_name).one()
    business = Business(name=name, sector_id=sector.id, region_id=region.id, google_place_id=f"gmaps_{name}", discovery_source="google_maps",
                        phone=phone, website=website, address="Test Cd. No:1", lat=40.77, lng=30.36, status="analyzing",
                        category_label=category, source_profile={"categories": [category]} if category else None)
    db.add(business)
    db.commit()
    job = AnalysisJob(business_id=business.id, user_id=user_id)
    db.add(job)
    db.commit()
    run_analysis_job(db, job.id)
    db.expire_all()
    return business.id


def test_detail_returns_opportunities_score_and_default_crm(client, db):
    bid = _analyzed_business(client, db)
    d = client.get(f"/api/businesses/{bid}").json()
    a = d["assessment"]
    assert d["business"]["crm_stage"] == "Yeni"
    assert set(a["opportunities"]) == {"evidence", "possible"}
    assert 0 <= a["score"]["score"] <= 100 and a["score"]["items"] is not None and a["score"]["note"]
    assert d["business"]["sales_score"] == a["score"]["score"]
    assert d["crm_history"] == []


def test_crm_stages_are_the_nine_sales_stages():
    assert CRM_STAGES == ["Yeni", "Aranacak", "Daha Sonra Ara", "Arandı", "Görüşüldü", "Teklif Gönderildi", "Takip Bekliyor", "Kazanıldı", "Kaybedildi"]


def test_crm_update_persists_and_logs_history(client, db):
    bid = _analyzed_business(client, db)
    r = client.patch(f"/api/businesses/{bid}/crm", json={"stage": "Aranacak", "staff_note": "Sahibi Salı 14:00'te aranacak"})
    assert r.status_code == 200 and r.json()["crm_stage"] == "Aranacak"
    again = client.get(f"/api/businesses/{bid}").json()
    assert again["business"]["crm_stage"] == "Aranacak" and again["business"]["staff_note"] == "Sahibi Salı 14:00'te aranacak"
    # CRM'de olmayan firmaya yapılan açık CRM işlemi firmayı CRM'e ekler: durum + not tek 'added' kaydıdır
    assert again["business"]["in_crm"] is True
    assert [h["type"] for h in again["crm_history"]] == ["added"]
    assert again["crm_history"][0]["to_stage"] == "Aranacak" and again["crm_history"][0]["note"] == "Sahibi Salı 14:00'te aranacak"
    listed = client.get("/api/businesses").json()
    assert next(b for b in listed if b["id"] == bid)["crm_stage"] == "Aranacak"


def test_crm_rejects_invalid_stage_and_empty_update(client, db):
    bid = _analyzed_business(client, db)
    assert client.patch(f"/api/businesses/{bid}/crm", json={"stage": "Uydurma"}).status_code == 422
    assert client.patch(f"/api/businesses/{bid}/crm", json={}).status_code == 422
    assert client.patch("/api/businesses/999999/crm", json={"stage": "Aranacak"}).status_code == 404


def test_all_nine_crm_stages_can_be_set(client, db):
    bid = _analyzed_business(client, db)
    for stage in CRM_STAGES:
        assert client.patch(f"/api/businesses/{bid}/crm", json={"stage": stage}).json()["crm_stage"] == stage


def test_guides_api_lists_library_and_renders_with_business_context(client, db):
    lib = client.get("/api/guides").json()
    assert set(lib["categories"]) == set(CATEGORIES) and len(lib["guides"]) >= 50
    bid = _analyzed_business(client, db, name="Örnek Merkezi")
    g = client.get(f"/api/guides/web_yok?business_id={bid}").json()
    assert set(g["sections"]) == {"what", "why", "solution", "steps", "pitch", "offer", "questions", "verify"}
    assert "Örnek Merkezi" in g["sections"]["pitch"] and g["business_evidence"], "işletmede gerçekten tespit edilen kanıt gösterilmeli"
    assert client.get("/api/guides/yok_boyle_rehber").status_code == 404


def test_sales_note_is_built_from_real_analysis(client, db):
    bid = _analyzed_business(client, db, name="Not Test Merkezi")
    note = client.get(f"/api/businesses/{bid}/sales-note").json()
    assert note["business"] == "Not Test Merkezi" and note["problems"]
    assert "TESPİT EDİLEN PROBLEMLER" in note["text"] and "ÖNERİLEN HİZMET" in note["text"] and "Web sitesi bulunamadı" in note["text"]
    assert note["questions"], "görüşmede sorulabilecek sorular olmalı"
    assert client.get("/api/businesses/999999/sales-note").status_code == 404


def test_sales_note_requires_analysis(client, db):
    region = db.query(Region).filter_by(name="Serdivan").one()
    sector = db.query(Sector).filter_by(name="İşitme Cihazı Merkezi").one()
    b = Business(name="Analizsiz", sector_id=sector.id, region_id=region.id, google_place_id="gmaps_x", discovery_source="google_maps")
    db.add(b)
    db.commit()
    assert client.get(f"/api/businesses/{b.id}/sales-note").status_code == 409


def test_csv_export_has_required_columns_and_real_values(client, db):
    bid = _analyzed_business(client, db, name="Dışa Aktarma Merkezi", phone="0532 999 88 77")
    client.patch(f"/api/businesses/{bid}/crm", json={"stage": "Takip Bekliyor", "staff_note": "=HYPERLINK(\"x\")"})
    r = client.post("/api/businesses/export", json={"ids": [bid], "format": "csv"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"] and r.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == ["İşletme adı", "Sektör", "İl", "İlçe", "Adres", "Telefon", "Web sitesi", "Google puanı", "Yorum sayısı", "Google profil eksikleri",
                       "Web sitesi eksikleri", "Sosyal medya durumu", "Satış fırsatı puanı", "Önerilen hizmetler", "CRM durumu", "Son analiz tarihi", "Personel notu"]
    row = dict(zip(rows[0], rows[1]))
    assert row["İşletme adı"] == "Dışa Aktarma Merkezi" and row["Telefon"] == "0532 999 88 77" and row["İlçe"] == "Serdivan" and row["İl"] == "Sakarya"
    assert row["Web sitesi"] == "Doğrulanamadı" and row["Google puanı"] == "Doğrulanamadı", "olmayan değer uydurulmamalı"
    assert row["CRM durumu"] == "Takip Bekliyor" and re.match(r"\d{2}\.\d{2}\.\d{4} \d{2}:\d{2}", row["Son analiz tarihi"])
    assert row["Personel notu"].startswith("'="), "formül enjeksiyonu engellenmeli"
    assert "Kurumsal Web Sitesi" in row["Önerilen hizmetler"]


def test_xlsx_export_opens_with_openpyxl(client, db):
    from openpyxl import load_workbook

    bid = _analyzed_business(client, db, name="Excel Merkezi")
    r = client.post("/api/businesses/export", json={"ids": [bid], "format": "xlsx"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    sheet = load_workbook(io.BytesIO(r.content)).active
    assert sheet["A1"].value == "İşletme adı" and sheet["A2"].value == "Excel Merkezi" and sheet.max_column == 17


def test_export_rejects_bad_requests(client):
    assert client.post("/api/businesses/export", json={"ids": [], "format": "csv"}).status_code == 422
    assert client.post("/api/businesses/export", json={"ids": [1], "format": "pdf"}).status_code == 422


def test_dashboard_today_lists_high_score_callable_businesses_only(client, db):
    hi = _analyzed_business(client, db, name="Yüksek Fırsat A.Ş.")
    lo = _analyzed_business(client, db, name="Kapanmış İş")
    client.patch(f"/api/businesses/{lo}/crm", json={"stage": "Kazanıldı"})
    today = client.get("/api/dashboard/today?min_score=1").json()
    ids = [b["id"] for b in today]
    assert hi in ids and lo not in ids, "müşteri olmuş/kapanmış işletmeler 'bugün ara' listesinde olmamalı"
    first = next(b for b in today if b["id"] == hi)
    assert first["phone"] and first["top_problem"] and first["primary_service"] and first["sales_score"] is not None
    scores = [b["sales_score"] for b in today]
    assert scores == sorted(scores, reverse=True)


def test_previously_analyzed_flag_relative_to_search_start(client, db):
    from packages.db.models import DiscoveryJob, DiscoveryJobResult

    me = client.get("/api/auth/me").json()["id"]
    bid = _analyzed_business(client, db, name="Eski Analiz Merkezi", user_id=me)
    region = db.query(Region).filter_by(name="Serdivan").one()
    sector = db.query(Sector).filter_by(name="İşitme Cihazı Merkezi").one()
    job = DiscoveryJob(region_id=region.id, sector_id=sector.id, target_count=5, status="completed")
    db.add(job)
    db.commit()
    db.add(DiscoveryJobResult(job_id=job.id, business_id=bid))
    db.commit()
    listed = client.get(f"/api/businesses?job_id={job.id}").json()
    assert listed[0]["previously_analyzed"] is True and listed[0]["last_analysis_at"], "iş başlamadan önce analiz edilen işletme 'daha önce analiz edildi' olmalı"


def test_dashboard_excludes_competitors_public_bodies_and_self_but_keeps_real_customers(client, db):
    keep = _analyzed_business(client, db, name="Özel Ölçün Anadolu Lisesi", category="Eğitim Kurumu", sector_name="Özel Okul")
    keep2 = _analyzed_business(client, db, name="Teknik Ses İşitme Cihazları", category="İşitme Cihazları Satıcısı")
    rival = _analyzed_business(client, db, name="Sakarya Web Tasarım Ajansı", category="Web Sitesi Tasarımcısı", sector_name="Web Tasarım ve Yazılım")
    public_body = _analyzed_business(client, db, name="Adapazarı Bilim ve Sanat Merkezi Müdürlüğü", category="Eğitim Kurumu", sector_name="Özel Okul")
    own = _analyzed_business(client, db, name="MchTTasarıM Reklam Ajansı | Web Tasarım", category="Web Sitesi Tasarımcısı", sector_name="Web Tasarım ve Yazılım")

    today = {b["id"] for b in client.get("/api/dashboard/today?min_score=1&limit=50").json()}
    assert keep in today and keep2 in today, "gerçek müşteri olabilecek işletmeler yanlışlıkla elenmemeli"
    assert not ({rival, public_body, own} & today), "rakip / kamu kurumu / Mchttasarım'ın kendisi panelde listelenmemeli"

    # ana arama listesinde hâlâ görünürler (yalnızca panelden çıkarılır), gerekçe alanı doludur
    listed = {b["id"]: b for b in client.get(f"/api/businesses?region_id={db.query(Region).filter_by(name='Serdivan').one().id}").json()}
    assert listed[rival]["prospect_kind"] == "rakip" and listed[rival]["prospect_reason"]
    assert listed[public_body]["prospect_kind"] == "kamu"
    assert listed[own]["prospect_kind"] == "kendi"
    assert listed[keep]["prospect_kind"] is None


def test_dashboard_still_fills_limit_after_excluding_top_scored_rivals(client, db):
    for i in range(4):
        _analyzed_business(client, db, name=f"Rakip Ajans {i}", category="Reklam Ajansı", sector_name="Reklam Ajansı")
    real = _analyzed_business(client, db, name="Gerçek Müşteri Kliniği")
    today = client.get("/api/dashboard/today?min_score=1&limit=1").json()
    assert [b["id"] for b in today] == [real], "elenenler yüzünden panel boş kalmamalı; sıradaki uygun aday listelenmeli"
