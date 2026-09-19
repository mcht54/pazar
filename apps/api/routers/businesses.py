from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Request, Response
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from apps.api.deps import client_ip, require
from apps.api.schemas import (
    AnalysisJobOut,
    AnalyzeBulkRequest,
    AssessmentOut,
    BusinessDetailOut,
    BusinessMetricOut,
    BusinessOut,
    CompetitorMetricRow,
    CrmAdd,
    CrmUpdate,
    ExportRequest,
    FindingOut,
    ServiceRecommendationOut,
)
from packages.db.base import get_db
from packages.db.models import (
    UserAnalysis,
    AnalysisJob,
    Business,
    DiscoveryJob,
    BusinessMetric,
    CompetitorSnapshot,
    DiscoveryJobResult,
    Finding,
    Region,
    SalesAssessment,
    Sector,
    ServiceCatalog,
    ServiceRecommendation,
    User,
)
from packages.crm import CRM_STAGES, normalize_stage
from services import crm_service
from services.auth.activity import log_activity
from services.auth.permissions import can
from services.knowledge.guides import guide_for_check
from services.rule_engine.competitor import COMPETITOR_METRIC_KEYS
from services.rule_engine.prospect import assess_business
from services.rule_engine.sales import LEVEL_ORDER, maps_search_url
from services.worker.tasks.analysis import SOURCE_LABELS, run_analysis_job_task

_TZ = ZoneInfo("Europe/Istanbul")

router = APIRouter(prefix="/api/businesses", tags=["işletmeler"])


def _latest_assessments(db: Session, business_ids: list[int]) -> dict[int, SalesAssessment]:
    if not business_ids:
        return {}
    latest = (
        select(func.max(SalesAssessment.id))
        .where(SalesAssessment.business_id.in_(business_ids))
        .group_by(SalesAssessment.business_id)
    )
    rows = db.query(SalesAssessment).filter(SalesAssessment.id.in_(latest)).all()
    return {row.business_id: row for row in rows}


def _region_labels(db: Session) -> dict[int, str]:
    regions = {r.id: r for r in db.query(Region).all()}
    labels = {}
    for region in regions.values():
        parent = regions.get(region.parent_region_id) if region.parent_region_id else None
        labels[region.id] = f"{region.name}, {parent.name}" if parent else region.name
    return labels


ANALYSIS_STATE_LABELS = {
    "none": "Henüz analiz edilmedi",
    "running": "Analiz ediliyor",
    "completed": "Analiz tamamlandı",
    "partial": "Kısmen analiz edildi",
    "failed": "Analiz başarısız",
}


def _load_extras(db: Session, businesses: list[Business], viewer: User | None = None) -> dict:
    """Kart/liste çıktıları için toplu yardımcı veriler: kullanıcı adları ve her firmanın SON analiz işi (tek sorgu)."""
    ids = [b.id for b in businesses]
    user_ids = {uid for b in businesses for uid in (b.crm_added_by, b.crm_updated_by, b.crm_owner_id) if uid}
    latest_jobs: dict[int, AnalysisJob] = {}
    if ids:
        latest = select(func.max(AnalysisJob.id)).where(AnalysisJob.business_id.in_(ids)).group_by(AnalysisJob.business_id)
        for job in db.query(AnalysisJob).filter(AnalysisJob.id.in_(latest)).all():
            latest_jobs[job.business_id] = job
            if job.user_id:
                user_ids.add(job.user_id)
        last_ok = select(func.max(AnalysisJob.id)).where(AnalysisJob.business_id.in_(ids), AnalysisJob.status.in_(("completed", "partial"))).group_by(AnalysisJob.business_id)
        finished = {j.business_id: j for j in db.query(AnalysisJob).filter(AnalysisJob.id.in_(last_ok)).all()}
        user_ids |= {j.user_id for j in finished.values() if j.user_id}
    else:
        finished = {}
    users = {u.id: u.name for u in db.query(User.id, User.name).filter(User.id.in_(user_ids)).all()} if user_ids else {}
    # KULLANICI BAZLI analiz geçmişi: yalnızca BAKAN kullanıcının kendi başarılı analizleri (viewer verilmediyse hiçbir şey "daha önce analiz edildi" sayılmaz)
    mine: dict[int, datetime] = {}
    if viewer is not None and ids:
        mine = {r.business_id: r.updated_at for r in db.query(UserAnalysis).filter(UserAnalysis.user_id == viewer.id, UserAnalysis.business_id.in_(ids), UserAnalysis.analysis_type == "deep").all()}
    return {"users": users, "latest_jobs": latest_jobs, "finished_jobs": finished, "my_analyses": mine, "viewer_id": viewer.id if viewer else None}


def _analysis_state(business: Business, job: AnalysisJob | None) -> str:
    """Kullanıcıya gösterilen analiz durumu: ÖNCE en son analiz işi; eski sonuç yeni durumla karıştırılmaz."""
    if business.status == "analyzing" or (job is not None and job.status in ("pending", "running")):
        return "running"
    if business.status == "analysis_failed" or (job is not None and job.status == "failed"):
        return "failed"
    if job is not None and job.status == "partial":
        return "partial"
    if business.last_analysis_at is not None:
        return "completed"
    return "none"


def _build_business_out(
    business: Business,
    *,
    sector_names: dict[int, str],
    region_labels: dict[int, str],
    service_names: dict[int, str],
    assessment: SalesAssessment | None,
    analyzed_before: datetime | None = None,
    extras: dict | None = None,
) -> BusinessOut:
    out = BusinessOut.model_validate(business)
    if extras is not None:
        users, job, finished = extras["users"], extras["latest_jobs"].get(business.id), extras["finished_jobs"].get(business.id)
        out.analysis_state = _analysis_state(business, job)
        out.analysis_state_label = ANALYSIS_STATE_LABELS[out.analysis_state]
        out.last_analyzed_by_name = users.get(finished.user_id) if finished and finished.user_id else None
        if business.crm_added_at is not None:
            out.crm_added_by_name = users.get(business.crm_added_by, "Bilinmiyor (eski kayıt)")
            out.crm_updated_by_name = users.get(business.crm_updated_by, "Bilinmiyor (eski kayıt)")
        out.crm_last_action = business.crm_last_action
        if business.crm_added_at is not None:
            out.crm_owner_name = users.get(business.crm_owner_id) if business.crm_owner_id else None
            out.follow_up_state = crm_service.follow_up_state(business)
    # "Daha önce analiz edildi": YALNIZCA bakan kullanıcının bu işletme için kendi başarılı analizi varsa (ve bu aramanın başlangıcından önceyse). Başka bir kullanıcının
    # analizi ya da işletmenin sistemde bulunması bunu doğurmaz.
    if extras is not None:
        mine_at = extras.get("my_analyses", {}).get(business.id)
        out.analyzed_by_me = mine_at is not None
        out.analyzed_by_me_at = mine_at
        out.analyzed_by_others = mine_at is None and business.last_analysis_at is not None
        out.previously_analyzed = mine_at is not None and (mine_at < analyzed_before if analyzed_before else True)
    out.sector_name = sector_names.get(business.sector_id, "")
    out.region_label = region_labels.get(business.region_id, "")
    out.in_crm = business.crm_added_at is not None
    parts = [p.strip() for p in out.region_label.split(",")]  # "Adapazarı, Sakarya" (ilçe, il) ya da yalnızca "Sakarya" (il)
    out.province_name = parts[-1] if parts and parts[-1] else None
    out.district_name = parts[0] if len(parts) > 1 else None
    verdict = assess_business(business, out.sector_name)
    if not verdict.eligible:
        out.prospect_kind, out.prospect_label, out.prospect_reason = verdict.kind, verdict.label, verdict.reason
    out.source_label = SOURCE_LABELS.get(business.discovery_source, business.discovery_source)
    out.is_demo_data = business.discovery_source == "mock_demo"
    out.source_url = (business.source_profile or {}).get("source_url")
    place = out.region_label.split(",")[0]
    out.maps_search_url = maps_search_url(business.name, business.address, place)
    if assessment is not None:
        out.sales_level = assessment.level
        out.sales_reason = assessment.level_reason
        out.needs_verification = assessment.needs_verification
        out.primary_service = service_names.get(assessment.primary_service_id)
        out.secondary_service = service_names.get(assessment.secondary_service_id)
        out.top_opportunity = assessment.top_opportunity
        payload = assessment.payload or {}
        out.gap_summary = [g["value"] for g in payload.get("gaps", [])[:4]]
        out.top_problem = out.gap_summary[0] if out.gap_summary else None
        priority = payload.get("priority") or {}
        out.priority_score = priority.get("value")
        out.priority_reasons = priority.get("reasons", [])
        out.why_prospect = payload.get("why_prospect", [])
        matrix_items = (payload.get("service_matrix") or {}).get("items", [])
        out.top_guide_id = next((i.get("guide_id") for i in matrix_items if i["level"] in ("satis", "olasi") and i.get("guide_id")), None) or next(
            (g.id for g in (guide_for_check(x.get("area", "website"), x.get("key", "")) for x in payload.get("gaps", [])) if g), None)
        out.sales_score = (payload.get("score") or {}).get("score", business.opportunity_score_total)
        out.score_band = (payload.get("score") or {}).get("band")
        out.why_call = assessment.why_call
        out.sales_note = assessment.sales_note
        out.gbp_summary = _area_summary(payload.get("gbp", {}), gbp=True)
        out.web_summary = _area_summary(payload.get("website", {}), gbp=False)
        research = payload.get("research") or {}
        out.verification = research.get("verdicts", {})
        out.source_statuses = research.get("sources", [])
        out.social_accounts = research.get("social", [])
        out.web_quick = _quick_rows((payload.get("website") or {}).get("checks", []), _WEB_QUICK)
        out.gbp_quick = _quick_rows((payload.get("gbp") or {}).get("checks", []), _GBP_QUICK)
        out.talking_point = assessment.talking_point
    return out


_WEB_QUICK = [("https", "HTTPS"), ("title", "Title"), ("meta_description", "Meta description"), ("h1", "H1"), ("local_seo", "Yerel SEO"),
              ("mobile", "Mobil"), ("whatsapp", "WhatsApp"), ("maps_link", "Google Maps"), ("services_page", "Hizmet sayfaları")]
_GBP_QUICK = [("primary_category", "Kategori"), ("description", "Açıklama"), ("services_list", "Hizmetler"), ("last_photo", "Fotoğraflar"),
              ("review_count", "Yorum sayısı"), ("rating", "Puan"), ("last_review", "Son yorum"), ("review_replies", "Yorum yanıtı"), ("hours", "Çalışma saatleri")]


def _quick_rows(checks: list[dict], spec: list[tuple[str, str]]) -> list[dict]:
    """Kart üzerinde tek satırlık özet: VAR / YOK / ZAYIF / DOĞRULANAMADI (+ ölçülen değer)."""
    by_key = {c["key"]: c for c in checks}
    rows = []
    for key, label in spec:
        check = by_key.get(key)
        if check is None:
            continue
        if check["status"] == "unknown":
            state = "DOĞRULANAMADI"
        elif check["status"] == "ok":
            state = "VAR"
        else:
            weak = key in ("local_seo",) or (key == "title" and not check["value"].startswith("Title etiketi yok")) or (
                key == "meta_description" and not check["value"].startswith("Meta description yok")
            ) or key in ("review_count", "rating", "last_review", "review_replies", "last_photo")
            state = "ZAYIF" if weak else "YOK"
        rows.append({"label": label, "state": state, "value": check["value"]})
    return rows


def _area_summary(area: dict, *, gbp: bool) -> str:
    """Kart üzerinde gösterilen tek satırlık Google profili / web sitesi analiz özeti."""
    checks = area.get("checks", [])
    problems = sum(1 for c in checks if c["status"] == "problem")
    unknown = sum(1 for c in checks if c["status"] == "unknown")
    if gbp and not area.get("available", False):
        return area.get("status_text") or "Google İşletme Profili verisi doğrulanamadı."
    if not gbp and not area.get("analyzed", False):
        return area.get("status_text", "Web sitesi analizi yok.")
    parts = [f"{problems} eksik/sorun tespit edildi"]
    if unknown:
        parts.append(f"{unknown} alan doğrulanamadı")
    return "; ".join(parts) + "."


def _presentation_maps(db: Session) -> tuple[dict[int, str], dict[int, str], dict[int, str]]:
    sector_names = {row.id: row.name for row in db.query(Sector.id, Sector.name).all()}
    service_names = {row.id: row.service_name for row in db.query(ServiceCatalog.id, ServiceCatalog.service_name).all()}
    return sector_names, _region_labels(db), service_names


@router.get("", response_model=list[BusinessOut])
def list_businesses(
    region_id: int | None = None,
    sector_id: int | None = None,
    crm_stage: str | None = None,
    job_id: int | None = None,
    level: str | None = None,
    user: User = Depends(require("view_business")),
    db: Session = Depends(get_db),
):
    """İşletmeleri satış önceliğine göre sıralı listeler: Yüksek → Orta → Düşük → Belirsiz → henüz analiz edilmemiş."""
    query = db.query(Business)
    if job_id is not None:
        query = query.join(DiscoveryJobResult, DiscoveryJobResult.business_id == Business.id).filter(DiscoveryJobResult.job_id == job_id)
    if region_id is not None:
        query = query.filter(Business.region_id == region_id)
    if sector_id is not None:
        query = query.filter(Business.sector_id == sector_id)
    if crm_stage is not None:
        if normalize_stage(crm_stage) is None:
            raise HTTPException(status_code=422, detail=f"Geçersiz CRM durumu. Geçerli durumlar: {', '.join(CRM_STAGES)}")
        query = query.filter(Business.crm_stage == crm_stage.strip())
    businesses = query.all()
    cutoff = None
    if job_id is not None:
        job = db.get(DiscoveryJob, job_id)
        cutoff = job.created_at if job else None  # bu aramadan önce analiz edilmiş işletmeler "daha önce analiz edildi" sayılır

    assessments = _latest_assessments(db, [b.id for b in businesses])
    if level is not None:
        businesses = [b for b in businesses if b.id in assessments and assessments[b.id].level == level]

    sector_names, region_labels, service_names = _presentation_maps(db)

    def sort_key(b: Business):
        """Varsayılan sıralama = satış ÖNCELİĞİ (skor + fırsat sayısı + ticari anlamlılık + analiz tamamlanma); eski analizlerde skor."""
        a = assessments.get(b.id)
        if a is None:
            return (-1.0, 0, 0.0, b.id)  # henüz analiz edilmemişler en sonda
        pr = ((a.payload or {}).get("priority") or {}).get("value")
        return (pr if pr is not None else float(b.opportunity_score_total or 0), LEVEL_ORDER.get(a.level, 0), a.rank_score, -b.id)

    businesses.sort(key=sort_key, reverse=True)
    extras = _load_extras(db, businesses, user)
    return [
        _build_business_out(b, sector_names=sector_names, region_labels=region_labels, service_names=service_names,
                            assessment=assessments.get(b.id), analyzed_before=cutoff, extras=extras)
        for b in businesses
    ]


def _analysis_history(db: Session, business_id: int, limit: int = 50) -> list[dict]:
    """Her analiz işlemi ayrı satırdır: kim başlattı · ne zaman tamamlandı · durum. (Aynı firma tekrar analiz edilirse yeni satır.)"""
    rows = (
        db.query(AnalysisJob, User.name)
        .outerjoin(User, User.id == AnalysisJob.user_id)
        .filter(AnalysisJob.business_id == business_id)
        .order_by(AnalysisJob.id.desc())
        .limit(limit)
        .all()
    )
    return [
        {"id": j.id, "status": j.status, "trigger": j.trigger, "user_id": j.user_id, "user_name": name or ("Bakım (otomatik)" if j.trigger == "maintenance" else "Bilinmiyor (eski kayıt)"),
         "started_at": j.started_at.isoformat() if j.started_at else None, "completed_at": j.completed_at.isoformat() if j.completed_at else None}
        for j, name in rows
    ]


def _build_competitor_rows(db: Session, business: Business, latest_job_id: int | None) -> list[CompetitorMetricRow]:
    if latest_job_id is None:
        return []
    snapshots = db.query(CompetitorSnapshot).filter_by(business_id=business.id, analysis_job_id=latest_job_id).all()
    if not snapshots:
        return []

    self_metrics = {
        m.metric_key: m
        for m in db.query(BusinessMetric)
        .filter(
            BusinessMetric.business_id == business.id,
            BusinessMetric.metric_key.in_(COMPETITOR_METRIC_KEYS),
            BusinessMetric.analysis_job_id.is_(None),
        )
        .order_by(BusinessMetric.id.asc())
        .all()
    }

    competitor_names = {
        row.id: row.name
        for row in db.query(Business).filter(Business.id.in_({s.competitor_business_id for s in snapshots})).all()
    }

    rows: dict[str, CompetitorMetricRow] = {}
    for metric_key in COMPETITOR_METRIC_KEYS:
        self_metric = self_metrics.get(metric_key)
        rows[metric_key] = CompetitorMetricRow(
            metric_key=metric_key,
            business_value=self_metric.value if self_metric else None,
            business_status=self_metric.status if self_metric else "not_available",
            competitors=[],
        )

    for snapshot in snapshots:
        rows[snapshot.metric_key].competitors.append(
            {
                "competitor_id": snapshot.competitor_business_id,
                "competitor_name": competitor_names.get(snapshot.competitor_business_id, "?"),
                "value": snapshot.value,
                "status": snapshot.status,
            }
        )

    return list(rows.values())


@router.get("/{business_id}", response_model=BusinessDetailOut)
def get_business_detail(business_id: int, user: User = Depends(require("view_business")), db: Session = Depends(get_db)):
    business = db.get(Business, business_id)
    if business is None:
        raise HTTPException(status_code=404, detail="İşletme bulunamadı")

    latest_job = db.query(AnalysisJob).filter_by(business_id=business_id).order_by(AnalysisJob.id.desc()).first()
    # Bulgular/ölçümler/öneriler: sadece EN SON tamamlanmış analiz — eski koşular birikmesin.
    finished_job = (
        db.query(AnalysisJob)
        .filter(AnalysisJob.business_id == business_id, AnalysisJob.status.in_(("completed", "partial")))
        .order_by(AnalysisJob.id.desc())
        .first()
    )
    finished_job_id = finished_job.id if finished_job else None

    metrics = (
        db.query(BusinessMetric)
        .filter(
            BusinessMetric.business_id == business_id,
            (BusinessMetric.analysis_job_id.is_(None)) | (BusinessMetric.analysis_job_id == finished_job_id),
        )
        .order_by(BusinessMetric.id.asc())
        .all()
    )
    findings = (
        db.query(Finding).filter(Finding.business_id == business_id, Finding.analysis_job_id == finished_job_id).order_by(Finding.id.asc()).all()
        if finished_job_id else []
    )
    recs = (
        db.query(ServiceRecommendation, ServiceCatalog.service_name)
        .join(ServiceCatalog, ServiceCatalog.id == ServiceRecommendation.service_id)
        .filter(ServiceRecommendation.business_id == business_id, ServiceRecommendation.analysis_job_id == finished_job_id)
        .order_by(ServiceRecommendation.priority_rank.asc())
        .all()
        if finished_job_id else []
    )
    assessment_row = _latest_assessments(db, [business_id]).get(business_id)

    sector_names, region_labels, service_names = _presentation_maps(db)
    extras = _load_extras(db, [business], user)
    business_out = _build_business_out(
        business, sector_names=sector_names, region_labels=region_labels, service_names=service_names, assessment=assessment_row, extras=extras
    )
    analysis_history = _analysis_history(db, business_id)

    assessment_out = None
    if assessment_row is not None:
        payload = assessment_row.payload or {}
        assessment_out = AssessmentOut(
            level=assessment_row.level,
            level_reason=assessment_row.level_reason,
            needs_verification=assessment_row.needs_verification,
            primary_service=service_names.get(assessment_row.primary_service_id),
            secondary_service=service_names.get(assessment_row.secondary_service_id),
            top_opportunity=assessment_row.top_opportunity,
            why_call=assessment_row.why_call,
            talking_point=assessment_row.talking_point,
            sales_note=assessment_row.sales_note,
            gaps=_guided(payload.get("gaps", [])),
            strengths=payload.get("strengths", []),
            services=payload.get("services", []),
            possible_services=payload.get("possible_services", []),
            verification_steps=payload.get("verification_steps", []),
            website={**(payload.get("website") or {}), "checks": _guided((payload.get("website") or {}).get("checks", []))},
            gbp={**(payload.get("gbp") or {}), "checks": _guided((payload.get("gbp") or {}).get("checks", []))},
            social=_guided(payload.get("social", [])),
            opportunities=payload.get("opportunities", {}),
            score=payload.get("score", {}),
            service_matrix=payload.get("service_matrix", {}),
            priority=payload.get("priority", {}),
            why_prospect=payload.get("why_prospect", []),
            analyzed_at=assessment_row.created_at,
        )

    return BusinessDetailOut(
        business=business_out,
        assessment=assessment_out,
        metrics=[BusinessMetricOut.model_validate(m) for m in metrics],
        findings=[FindingOut.model_validate(f) for f in findings],
        service_recommendations=[
            ServiceRecommendationOut(
                service_id=rec.service_id, service_name=name, priority_rank=rec.priority_rank,
                matched_rule=rec.matched_rule, ai_justification=rec.ai_justification,
            )
            for rec, name in recs
        ],
        competitors=_build_competitor_rows(db, business, finished_job_id),
        latest_analysis_job=AnalysisJobOut.model_validate(latest_job) if latest_job else None,
        crm_history=crm_service.history(db, business_id),
        follow_ups=crm_service.follow_up_list(db, business_id),
        analysis_history=analysis_history,
    )


def _guided(checks: list[dict]) -> list[dict]:
    """Her kontrol/eksik satırına, varsa "💡 Nasıl çözülür?" rehberinin kimliğini ekler."""
    out = []
    for check in checks:
        guide = guide_for_check(check.get("area", "website"), check.get("key", ""))
        out.append({**check, "guide_id": guide.id if guide else None})
    return out


def _sales_fields_of(payload, user: User) -> dict:
    """İstekte gönderilen satış takibi alanları; sorumlu personel değişikliği ayrı yetki ister (Yönetici/Çalışan)."""
    fields = payload.provided()
    if "owner_id" in fields and not can(user.role, "crm_assign"):
        raise HTTPException(status_code=403, detail="Sorumlu personeli değiştirme yetkiniz yok.")
    return fields


def _crm_response(db: Session, business: Business) -> dict:
    return crm_service.crm_state(db, business)


@router.post("/{business_id}/crm")
def add_to_crm(business_id: int, payload: CrmAdd, request: Request, user: User = Depends(require("crm_use")), db: Session = Depends(get_db)):
    """'➕ CRM'e Ekle': firmayı, seçilen durum ve (isteğe bağlı) notla CRM'e alır. Zaten CRM'deyse 409 döner."""
    business = db.get(Business, business_id)
    if business is None:
        raise HTTPException(status_code=404, detail="İşletme bulunamadı")
    try:
        crm_service.add_to_crm(db, business, payload.stage, payload.note, user=user, ip=client_ip(request), fields=_sales_fields_of(payload, user))
    except crm_service.CrmError as exc:
        raise HTTPException(status_code=exc.status_code, detail=exc.detail) from exc
    return _crm_response(db, business)


@router.patch("/{business_id}/crm")
def update_crm(business_id: int, payload: CrmUpdate, request: Request, user: User = Depends(require("crm_use")), db: Session = Depends(get_db)):
    """CRM durumunu ve/veya CRM notunu KALICI olarak günceller (son işlem zamanı güncellenir, geçmişe kayıt yazılır).

    CRM'de olmayan firmaya yapılan bu çağrı açık bir CRM işlemi sayılır ve firmayı CRM'e ekler; analiz ise asla eklemez.
    """
    business = db.get(Business, business_id)
    if business is None:
        raise HTTPException(status_code=404, detail="İşletme bulunamadı")
    fields = _sales_fields_of(payload, user)
    if payload.stage is None and payload.staff_note is None and not fields:
        raise HTTPException(status_code=422, detail="Güncellenecek bir alan gönderilmedi (stage, staff_note, takip/teklif/satış alanları)")
    try:
        crm_service.update_crm(db, business, payload.stage, payload.staff_note, user=user, ip=client_ip(request), fields=fields)
    except crm_service.CrmError as exc:
        raise HTTPException(status_code=exc.status_code, detail=exc.detail) from exc
    return _crm_response(db, business)


@router.get("/{business_id}/sales-note")
def sales_note(business_id: int, request: Request, user: User = Depends(require("sales_note")), db: Session = Depends(get_db)):
    """'Satış Notu Oluştur': gerçek analiz sonuçlarından personel için kısa satış notu (uydurma bilgi içermez)."""
    from services.knowledge.sales_note import build_sales_note

    business = db.get(Business, business_id)
    if business is None:
        raise HTTPException(status_code=404, detail="İşletme bulunamadı")
    assessment = _latest_assessments(db, [business_id]).get(business_id)
    if assessment is None:
        raise HTTPException(status_code=409, detail="Satış notu için önce işletmenin analiz edilmesi gerekir")
    sector = db.get(Sector, business.sector_id)
    region_label = _region_labels(db).get(business.region_id, "")
    _, _, service_names = _presentation_maps(db)
    result = build_sales_note(
        business, assessment.payload or {}, sector=sector.name, place=region_label.split(",")[0],
        phrase=(sector.keyword_variants or [sector.name.lower()])[0], level=assessment.level,
        primary=service_names.get(assessment.primary_service_id), secondary=service_names.get(assessment.secondary_service_id),
        talking_point=assessment.talking_point, sales_pitch=assessment.sales_note,
    )
    log_activity(db, user, "sales_note", business_id=business_id, detail=f"{business.name}: satış notu oluşturuldu", ip=client_ip(request), commit=True)
    return result


_EXPORT_COLUMNS = [
    "İşletme adı", "Sektör", "İl", "İlçe", "Adres", "Telefon", "Web sitesi", "Google puanı", "Yorum sayısı", "Google profil eksikleri",
    "Web sitesi eksikleri", "Sosyal medya durumu", "Satış fırsatı puanı", "Önerilen hizmetler", "CRM durumu", "Son analiz tarihi", "Personel notu",
]
_UNVERIFIED = "Doğrulanamadı"


def _csv_safe(value) -> str:
    """CSV/Excel formül enjeksiyonunu önler (=, +, -, @ ile başlayan hücreler metin olarak işaretlenir)."""
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


def _export_rows(db: Session, ids: list[int]) -> list[list[str]]:
    businesses = db.query(Business).filter(Business.id.in_(ids)).all()
    order = {bid: i for i, bid in enumerate(ids)}
    businesses.sort(key=lambda b: order.get(b.id, 0))
    assessments = _latest_assessments(db, ids)
    sector_names, _, _ = _presentation_maps(db)
    regions = {r.id: r for r in db.query(Region).all()}
    rows = []
    for b in businesses:
        region = regions.get(b.region_id)
        parent = regions.get(region.parent_region_id) if region and region.parent_region_id else None
        il, ilce = (parent.name, region.name) if parent else (region.name if region else "", "")
        a = assessments.get(b.id)
        payload = (a.payload if a else None) or {}
        gbp, web = payload.get("gbp") or {}, payload.get("website") or {}
        gbp_missing = "; ".join(c["value"] for c in gbp.get("checks", []) if c["status"] == "problem") if gbp.get("available") else _UNVERIFIED
        web_missing = "; ".join(c["value"] for c in web.get("checks", []) if c["status"] == "problem") if a else _UNVERIFIED
        social = "; ".join(f"{s['label']}: {'doğrulandı' if s['status'] == 'dogrulandi' else 'bulunamadı' if s['status'] == 'bulunamadi' else 'doğrulanmadı'}"
                           for s in (payload.get("research") or {}).get("social", [])) or _UNVERIFIED
        services = ", ".join(x["service"] for x in (payload.get("opportunities") or {}).get("evidence", [])) or ("-" if a else _UNVERIFIED)
        rows.append([
            b.name, sector_names.get(b.sector_id, ""), il, ilce, b.address or _UNVERIFIED, b.phone or _UNVERIFIED, b.website or _UNVERIFIED,
            str(b.google_rating).replace(".", ",") if b.google_rating is not None else _UNVERIFIED,
            b.google_review_count if b.google_review_count is not None else _UNVERIFIED,
            gbp_missing or "Eksik tespit edilmedi", web_missing or "Eksik tespit edilmedi", social,
            (payload.get("score") or {}).get("score", b.opportunity_score_total) if a else _UNVERIFIED,
            services, b.crm_stage if b.crm_added_at else "CRM'de değil",
            b.last_analysis_at.astimezone(_TZ).strftime("%d.%m.%Y %H:%M") if b.last_analysis_at else "Analiz edilmedi", b.staff_note or "",
        ])
    return rows


@router.post("/export")
def export_businesses(payload: ExportRequest, request: Request, user: User = Depends(require("export")), db: Session = Depends(get_db)):
    """Seçili/görünen işletmeleri CSV (Excel'de doğrudan açılır: UTF-8 BOM + ';') veya XLSX olarak dışa aktarır."""
    if not payload.ids:
        raise HTTPException(status_code=422, detail="Dışa aktarılacak işletme yok")
    if payload.format not in ("csv", "xlsx"):
        raise HTTPException(status_code=422, detail="Biçim csv veya xlsx olmalı")
    rows = _export_rows(db, payload.ids[:1000])
    log_activity(db, user, "export", detail=f"{len(rows)} firma ({payload.format.upper()})", meta={"count": len(rows), "format": payload.format}, ip=client_ip(request), commit=True)
    stamp = datetime.now().strftime("%Y%m%d")
    if payload.format == "csv":
        import csv
        import io

        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        writer.writerow(_EXPORT_COLUMNS)
        writer.writerows([[_csv_safe(v) for v in row] for row in rows])
        return Response(content="\ufeff" + buffer.getvalue(), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="mchttasarim_isletmeler_{stamp}.csv"'})

    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "İşletmeler"
    sheet.append(_EXPORT_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A56DB")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        sheet.append([_csv_safe(v) if isinstance(v, str) else v for v in row])
    for column, width in zip("ABCDEFGHIJKLMNOPQ", [34, 22, 14, 16, 42, 16, 30, 12, 12, 46, 46, 34, 14, 40, 16, 16, 40]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "B2"
    output = io.BytesIO()
    workbook.save(output)
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="mchttasarim_isletmeler_{stamp}.xlsx"'})


@router.post("/{business_id}/analyze", response_model=AnalysisJobOut)
def analyze_business(business_id: int, request: Request, user: User = Depends(require("analyze")), db: Session = Depends(get_db)):
    business = db.get(Business, business_id)
    if business is None:
        raise HTTPException(status_code=404, detail="İşletme bulunamadı")

    job = AnalysisJob(business_id=business_id, user_id=user.id)  # analiz geçmişi: kim başlattı
    db.add(job)
    business.status = "analyzing"
    db.flush()
    log_activity(db, user, "analyze", business_id=business_id, detail=f"{business.name}: analiz başlatıldı", meta={"analysis_job_id": job.id}, ip=client_ip(request))
    db.commit()
    db.refresh(job)

    run_analysis_job_task.delay(job.id, True)  # kullanıcı isteği: kaynaklar yeniden sorgulanır
    return job


@router.post("/analyze-bulk", response_model=list[AnalysisJobOut])
def analyze_bulk(payload: AnalyzeBulkRequest, request: Request, user: User = Depends(require("analyze")), db: Session = Depends(get_db)):
    if not payload.business_ids and not payload.top_n:
        raise HTTPException(status_code=422, detail="İşletme seçimi (business_ids) veya top_n belirtilmeli — tüm işletmeler otomatik analiz edilmez")

    if payload.business_ids:
        businesses = db.query(Business).filter(Business.id.in_(payload.business_ids)).all()
    else:
        if payload.top_n <= 0 or payload.top_n > 100:
            raise HTTPException(status_code=422, detail="top_n 1-100 arasında olmalı")
        businesses = (
            db.query(Business)
            .filter(Business.status.in_(("discovered", "analysis_failed")))
            .order_by(Business.created_at.asc())
            .limit(payload.top_n)
            .all()
        )

    jobs = []
    for business in businesses:
        job = AnalysisJob(business_id=business.id, user_id=user.id)
        db.add(job)
        business.status = "analyzing"
        db.flush()
        log_activity(db, user, "analyze", business_id=business.id, detail=f"{business.name}: analiz başlatıldı (toplu)", meta={"analysis_job_id": job.id}, ip=client_ip(request))
        jobs.append(job)
    db.commit()

    for job in jobs:
        db.refresh(job)
        run_analysis_job_task.delay(job.id, True)

    return jobs
